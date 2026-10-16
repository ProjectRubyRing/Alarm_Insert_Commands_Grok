import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "CloudWatch_Alarm_安全ログ追記コマンド集.xlsx"
BASH = Path(r"C:\Program Files\Git\bin\bash.exe")


def gray(value) -> bool:
    if value is None:
        return True
    text = str(value)
    if text in {"0", "1"}:
        return True
    if len(text) == 8:
        return text[2:4] == text[4:6] == text[6:8]
    return False


class WorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.book = load_workbook(BOOK)

    def test_sheets_font_and_monotone(self):
        self.assertEqual(
            self.book.sheetnames,
            ["はじめに", "コマンド一覧", "くわしい説明", "実行チェック", "事故を防ぐ", "言葉と年表"],
        )
        for ws in self.book.worksheets:
            self.assertFalse(ws.sheet_view.showGridLines)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    self.assertEqual(cell.font.name, "Meiryo UI", cell.coordinate)
                    self.assertTrue(cell.alignment.wrap_text, cell.coordinate)
                    self.assertEqual(cell.number_format, "@", cell.coordinate)
                    self.assertTrue(gray(cell.font.color.rgb if cell.font.color is not None else None), cell.coordinate)
                    if cell.fill.fgColor is not None:
                        self.assertTrue(gray(cell.fill.fgColor.rgb), f"{ws.title} {cell.coordinate} {cell.fill.fgColor.rgb}")

    def test_commands_are_safe_and_parse(self):
        ws = self.book["コマンド一覧"]
        headers = [cell.value for cell in ws[4]]
        step1 = headers.index("手順1_先にコピー") + 1
        step2 = headers.index("手順2_確認後にコピー") + 1
        id_col = headers.index("ID") + 1
        bash_count = 0
        ps_count = 0
        for row in range(5, ws.max_row + 1):
            command_id = ws.cell(row, id_col).value
            for col in (step1, step2):
                text = ws.cell(row, col).value or ""
                self.assertNotIn("delete-log-group", text)
                self.assertNotIn("delete-log-stream", text)
                self.assertNotIn("--sequence-token", text)
                self.assertNotIn("--no-deletion-protection-enabled", text)
                for line in text.splitlines():
                    if line.strip().startswith("#"):
                        continue
                    self.assertNotIn("Set-Content", line)
                    self.assertNotIn("Clear-Content", line)
                    stripped = line.replace(">>", "")
                    if ">" in stripped and ("LOG_FILE" in stripped or "$LogFile" in stripped):
                        self.fail(f"{command_id} overwrites the log: {line}")
                if "$LogFile" in text or "$payload" in text:
                    ps_count += 1
                    self.assert_powershell(command_id, text)
                elif text.startswith("#") or text.startswith("この手順"):
                    if text.startswith("この手順"):
                        continue
                    bash_count += 1
                    self.assert_bash(command_id, text)
        self.assertGreater(bash_count, 10)
        self.assertGreaterEqual(ps_count, 2)

    def assert_bash(self, command_id, text):
        with tempfile.NamedTemporaryFile("w", suffix=".sh", delete=False, encoding="utf-8", newline="\n") as handle:
            handle.write(text)
            path = handle.name
        result = subprocess.run([str(BASH), "-n", path], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, f"{command_id}\n{result.stderr}\n{text}")

    def assert_powershell(self, command_id, text):
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig", newline="\r\n") as handle:
            handle.write(text)
            path = handle.name
        checker = (
            "$errs=$null; "
            f"[void][System.Management.Automation.Language.Parser]::ParseFile('{path}', [ref]$null, [ref]$errs); "
            "if ($errs) { $errs | ForEach-Object { $_.ToString() }; exit 1 } else { exit 0 }"
        )
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-Command", checker],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        self.assertEqual(result.returncode, 0, f"{command_id}\n{result.stdout}\n{result.stderr}\n{text}")


if __name__ == "__main__":
    unittest.main()
